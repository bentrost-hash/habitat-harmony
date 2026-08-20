"""Normalize a downloaded Planet Zoo Data Chest workbook into app data.

Usage (after downloading the public workbook):
  python tools/update_data.py source-data.xlsx data/animals.json

The app never reads the workbook; this is a deliberate, offline build step.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl


def clean(value):
    if value is None or value == "-":
        return None
    return str(value).strip()


def words(value):
    return [clean(value)] if clean(value) else []


def identifier(name):
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def temperature(value):
    # Zoopedia uses both “3 - 45” and “13-38”; hyphens are range separators.
    text = str(value or "").strip()
    numbers = re.findall(r"-?\d+(?:\.\d+)?", text)
    if not text.startswith("-"):
        numbers = [item.lstrip("-") for item in numbers]
    if len(numbers) != 2:
        return None
    return {"min": float(numbers[0]), "max": float(numbers[1])}


def main(source, destination):
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    habitat = book["Habitat Animals Data"]
    animals = []
    by_name = {}
    for row in habitat.iter_rows(min_row=3, values_only=True):
        name = clean(row[0])
        if not name:
            continue
        animal = {
            "id": identifier(name), "name": name,
            "dlc": clean(row[30]) or "Unknown",
            "continents": [x for x in map(clean, row[1:5]) if x],
            "biomes": [x for x in map(clean, row[5:10]) if x],
            "temperature": temperature(row[14]),
            "landRequirement": row[11] if isinstance(row[11], (int, float)) else None,
            "familyLandRequirement": row[12] if isinstance(row[12], (int, float)) else None,
            "barrier": {"height": row[15], "grade": row[16]},
            "climbing": clean(row[23]) == "Y",
            "water": {"canSwim": clean(row[24]) == "Y", "deepDiver": clean(row[25]) == "Y"},
            "walkthroughSafe": clean(row[18]) == "Y",
            "aggressive": clean(row[28]) == "Aggressive",
            "interspeciesEnrichment": [],
            "compatibility": {},
            "source": "Villanelle Planet Zoo Data Chest (final 2025 workbook)"
        }
        animals.append(animal)
        by_name[name] = animal

    chart = book["Compatibility Chart"]
    headers = [clean(cell.value) for cell in chart[1]][1:]
    raw = {}
    source_asymmetries = []
    for row in chart.iter_rows(min_row=2, max_row=len(headers) + 1, values_only=True):
        name = clean(row[0])
        if name not in by_name:
            continue
        animal = by_name[name]
        for partner, rating in zip(headers, row[1:]):
            if not partner or partner == name or partner not in by_name or rating is None:
                continue
            raw[(animal["id"], by_name[partner]["id"])] = 4 if rating == "⭐" else int(rating)

    # The chart contains directional predator/partner entries in a few places.
    # A shared habitat must be safe in both directions, so retain the lower score.
    for left in animals:
        for right in animals:
            if left["id"] >= right["id"]: continue
            a, b = left["id"], right["id"]
            first, second = raw.get((a, b)), raw.get((b, a))
            if first is None or second is None: continue
            if first != second:
                source_asymmetries.append({"pair": [a, b], "ratings": [first, second]})
            score = min(first, second)
            left["compatibility"][b] = score
            right["compatibility"][a] = score
            if first == 4 or second == 4:
                left["interspeciesEnrichment"].append(b)
                right["interspeciesEnrichment"].append(a)

    payload = {"schemaVersion": 1, "generatedFrom": Path(source).name,
               "animalCount": len(animals), "animals": animals}
    Path(destination).parent.mkdir(parents=True, exist_ok=True)
    Path(destination).write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    Path(destination).with_suffix(".js").write_text(
        "window.PZ_DATA = " + json.dumps(payload, ensure_ascii=False) + ";\n", encoding="utf-8")
    report = {"sourceDirectionalPairs": source_asymmetries,
              "note": "Canonical pair score uses the more restrictive source rating."}
    Path(destination).with_name("normalization-report.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8")
    print(f"Wrote {len(animals)} habitat animals to {destination}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: update_data.py SOURCE.xlsx data/animals.json")
    main(sys.argv[1], sys.argv[2])
